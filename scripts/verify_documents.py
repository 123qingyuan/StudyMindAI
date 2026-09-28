"""Local document engine acceptance. Generated files are labelled test fixtures."""
from pathlib import Path
import json
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from PIL import Image,ImageDraw,ImageFont
import pymupdf
from docx import Document
from app.document.parser import extract

root=Path(__file__).resolve().parents[1]
folder=root/'data'/'acceptance-fixtures'
folder.mkdir(parents=True,exist_ok=True)
phrase='数据库主键用于唯一标识一条记录'
image=Image.new('RGB',(1400,220),'white')
draw=ImageDraw.Draw(image)
font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',44)
draw.text((30,65),phrase,fill='black',font=font)
image.save(folder/'ocr-test.png')
page_pdf=pymupdf.open();page=page_pdf.new_page();page.insert_text((50,80),'Database primary key uniquely identifies a record.',fontsize=15);page_pdf.save(folder/'text-test.pdf');page_pdf.close()
scan=pymupdf.open();page=scan.new_page(width=700,height=110);page.insert_image(page.rect,filename=str(folder/'ocr-test.png'));scan.save(folder/'scan-test.pdf');scan.close()
document=Document();document.add_heading('数据库测试资料',0);document.add_paragraph(phrase);document.save(folder/'test.docx')
(folder/'test.txt').write_text(phrase,encoding='utf-8')
results=[]
for name in ['test.txt','test.docx','text-test.pdf','ocr-test.png','scan-test.pdf']:
    pages=extract(folder/name,Path(name).suffix)
    content='\n'.join(p.text for p in pages)
    expected='primary key' if name=='text-test.pdf' else '主键'
    assert expected in content,(name,content)
    results.append({'file':name,'text':content,'methods':list({p.method for p in pages}),'ok':True})
(root/'data'/'document-acceptance.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=False))
