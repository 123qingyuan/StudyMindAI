"""Bounded real document parsing. OCR uses packaged Chinese ONNX models."""
from __future__ import annotations
import io
import threading
import zipfile
from dataclasses import dataclass
from pathlib import Path
from PIL import Image, UnidentifiedImageError

MAX_BYTES = 20 * 1024 * 1024
MAX_PAGES = 100
MAX_TEXT_CHARS = 600_000
MAX_IMAGE_PIXELS = 20_000_000
ALLOWED = {'.txt', '.md', '.markdown', '.csv', '.json', '.docx', '.pdf', '.png', '.jpg', '.jpeg', '.webp'}
MIMES = {'.pdf': 'application/pdf', '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.webp': 'image/webp'}
_ocr = None
_ocr_lock = threading.Lock()


class DocumentError(Exception):
    def __init__(self, code, message, status=422):
        self.code, self.message, self.status = code, message, status
        super().__init__(code)


@dataclass
class Page:
    number: int | None
    text: str
    method: str


def safe_display_name(name):
    name = (name or '未命名').replace('\\', '/').split('/')[-1]
    return ''.join(c for c in name if ord(c) >= 32 and ord(c) != 127)[:200] or '未命名'


def validate_docx(path):
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                raise DocumentError('DOCX_ZIP_LIMIT', 'DOCX 解压后大小或文件数量超限')
            for info in infos:
                if info.flag_bits & 1 or info.file_size > 30 * 1024 * 1024 or info.file_size > max(info.compress_size, 1) * 200:
                    raise DocumentError('DOCX_ZIP_LIMIT', 'DOCX 包含过度压缩或加密内容')
                if info.filename.startswith('/') or '..' in Path(info.filename).parts or '\\' in info.filename:
                    raise DocumentError('INVALID_DOCX', 'DOCX 包含非法路径')
            names = archive.namelist()
            if len(set(names)) != len(names) or 'word/document.xml' not in names or '[Content_Types].xml' not in names:
                raise DocumentError('INVALID_DOCX', '文件不是有效 DOCX')
            if any(name.lower().endswith('vbaproject.bin') for name in names):
                raise DocumentError('INVALID_DOCX', '不支持含宏的文档')
    except zipfile.BadZipFile:
        raise DocumentError('INVALID_DOCX', 'DOCX 压缩结构损坏') from None


def validate_signature(path, extension):
    with open(path, 'rb') as source:
        prefix = source.read(32)
    valid = True
    if extension == '.pdf':
        valid = prefix.startswith(b'%PDF-')
    elif extension == '.docx':
        valid = prefix.startswith(b'PK\x03\x04')
    elif extension == '.png':
        valid = prefix.startswith(b'\x89PNG\r\n\x1a\n')
    elif extension in {'.jpg', '.jpeg'}:
        valid = prefix.startswith(b'\xff\xd8\xff')
    elif extension == '.webp':
        valid = prefix.startswith(b'RIFF') and prefix[8:12] == b'WEBP'
    if not valid:
        raise DocumentError('FILE_SIGNATURE_MISMATCH', '文件内容与扩展名不一致', 415)
    if extension == '.docx':
        validate_docx(path)


def ocr_image(image: Image.Image) -> str:
    global _ocr
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise DocumentError('IMAGE_TOO_LARGE', '图片像素数量超出限制')
    try:
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        raise DocumentError('OCR_UNAVAILABLE', 'OCR 组件未安装，请联系管理员', 503) from None
    with _ocr_lock:
        if _ocr is None:
            _ocr = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=2,
                            det_use_cuda=False, cls_use_cuda=False, rec_use_cuda=False)
        # RapidOCR ndarray input uses OpenCV BGR channel order.
        array = np.asarray(image.convert('RGB'))[:, :, ::-1].copy()
        result, _ = _ocr(array)
    return '\n'.join(str(row[1]) for row in (result or []) if row[1] and float(row[2]) >= 0.45)


def image_page(path, number=1):
    try:
        with Image.open(path) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS or getattr(image, 'n_frames', 1) > 1:
                raise DocumentError('IMAGE_TOO_LARGE', '图片像素过大或为多帧图片')
            image.load()
            return Page(number, ocr_image(image), 'ocr')
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError):
        raise DocumentError('INVALID_IMAGE', '图片损坏或无法解析') from None


def extract(path: Path, extension: str) -> list[Page]:
    if extension not in ALLOWED:
        raise DocumentError('UNSUPPORTED_FILE', '文件格式不受支持', 415)
    validate_signature(path, extension)
    if extension in {'.png', '.jpg', '.jpeg', '.webp'}:
        pages = [image_page(path)]
    elif extension == '.pdf':
        import pymupdf
        try:
            with pymupdf.open(path) as pdf:
                if pdf.needs_pass:
                    raise DocumentError('ENCRYPTED_PDF', '暂不支持加密 PDF')
                if not 1 <= len(pdf) <= MAX_PAGES:
                    raise DocumentError('PDF_PAGE_LIMIT', f'PDF 最多允许 {MAX_PAGES} 页')
                pages = []
                for index, page in enumerate(pdf):
                    text = page.get_text('text').strip()
                    if not text or (len(text) < 40 and page.get_images()):
                        # Avoid huge rasterization before checking its allocation.
                        rect = page.rect
                        scale = min(2.0, (MAX_IMAGE_PIXELS / max(rect.width * rect.height, 1)) ** 0.5)
                        if scale < 0.25:
                            raise DocumentError('PDF_PAGE_TOO_LARGE', 'PDF 页面尺寸异常')
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                        image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
                        scanned = ocr_image(image)
                        text = scanned if scanned else text
                        method = 'ocr' if scanned else 'text'
                    else:
                        method = 'text'
                    pages.append(Page(index + 1, text, method))
                    if sum(len(p.text) for p in pages) > MAX_TEXT_CHARS:
                        raise DocumentError('TEXT_LIMIT', '提取文本超过安全限制，请拆分资料')
        except DocumentError:
            raise
        except Exception:
            raise DocumentError('INVALID_PDF', 'PDF 损坏或解析失败') from None
    elif extension == '.docx':
        from docx import Document
        try:
            doc = Document(str(path))
            texts = []
            # iter_inner_content preserves paragraph/table ordering (python-docx 1.2+).
            from docx.table import Table
            for block in doc.iter_inner_content():
                if isinstance(block, Table):
                    texts.extend('\t'.join(cell.text for cell in row.cells) for row in block.rows)
                else:
                    texts.append(block.text)
            for section in doc.sections:
                texts.extend(p.text for p in section.header.paragraphs)
                texts.extend(p.text for p in section.footer.paragraphs)
            # DOCX pagination cannot be determined without a layout engine.
            pages = [Page(None, '\n'.join(texts), 'docx')]
        except Exception:
            raise DocumentError('INVALID_DOCX', 'DOCX 解析失败') from None
    else:
        data = path.read_bytes()
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = data.decode('gb18030')
            except UnicodeDecodeError:
                raise DocumentError('INVALID_TEXT_ENCODING', '文本必须是 UTF-8 或 GB18030 编码') from None
        if '\x00' in text or sum(ord(c) < 32 and c not in '\r\n\t' for c in text) > 10:
            raise DocumentError('INVALID_TEXT_FILE', '文本文件包含二进制内容', 415)
        pages = [Page(None, text, 'text')]
    if sum(len(page.text) for page in pages) > MAX_TEXT_CHARS:
        raise DocumentError('TEXT_LIMIT', '提取文本超过安全限制，请拆分资料')
    if not any(page.text.strip() for page in pages):
        raise DocumentError('NO_EXTRACTABLE_TEXT', '没有提取到可用文字；OCR 未识别到清晰文本')
    return pages


def chunks(pages: list[Page], size=900, overlap=120):
    ordinal = 0
    for page in pages:
        text = page.text.strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            yield {'page': page.number, 'ordinal': ordinal, 'content': text[start:end]}
            ordinal += 1
            if end == len(text):
                break
            start = end - overlap
